"""Generate Baltan pitch deck (PPTX) for PRARAMBH / incubator audience.

Plain-language, visually strong. Funding slide = why + use of funds (no ask amount).

Run: python generate_baltan_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Baltan brand
BG = RGBColor(0x07, 0x0A, 0x0C)
ELEV = RGBColor(0x10, 0x16, 0x1A)
ELEV2 = RGBColor(0x16, 0x1E, 0x24)
LINE = RGBColor(0x24, 0x30, 0x38)
TEXT = RGBColor(0xEE, 0xF6, 0xF3)
MUTED = RGBColor(0x8A, 0xA0, 0x99)
ACCENT = RGBColor(0x1D, 0xE9, 0xB6)
ACCENT_DIM = RGBColor(0x0F, 0x3D, 0x34)
WARN = RGBColor(0xE0, 0xB3, 0x4A)
SOFT = RGBColor(0x1A, 0x2A, 0x32)

ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "apps" / "web" / "public" / "brand"
LOGO = BRAND / "baltan-logo-dark-bg-2x.png"
ICON = BRAND / "baltan-icon-256.png"
OUT_DIR = Path(__file__).resolve().parent
OUT_PPTX = OUT_DIR / "Baltan_Pitch_Deck.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 9


def set_run(run, text, size=18, bold=False, color=TEXT, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=TEXT,
    align=PP_ALIGN.LEFT,
    font="Calibri",
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        r = p.add_run()
        set_run(r, line, size=size, bold=bold, color=color, font=font)
    return box


def fill_shape(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def oval(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    fill_shape(sh, fill)
    return sh


def rect(slide, left, top, width, height, fill=ELEV, line=None, radius=0.1):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(sh, fill)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    if line is not None:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    return sh


def accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.55)):
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(bar, ACCENT)
    return bar


def chip(slide, left, top, width, height, text, bg=ACCENT_DIM, fg=ACCENT, size=11):
    sh = rect(slide, left, top, width, height, fill=bg, line=ACCENT, radius=0.35)
    tf = sh.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, text, size=size, bold=True, color=fg)
    return sh


def set_slide_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    fill_shape(bg, BG)
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    # soft decorative glow (top-right)
    oval(slide, Inches(10.2), Inches(-1.2), Inches(4.5), Inches(4.5), SOFT)
    oval(slide, Inches(-1.5), Inches(5.2), Inches(3.2), Inches(3.2), ACCENT_DIM)


def footer(slide, n):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.02), Inches(12.2), Pt(1.5)
    )
    fill_shape(line, LINE)
    add_textbox(slide, Inches(0.55), Inches(7.12), Inches(3), Inches(0.28), "baltan", size=11, color=MUTED)
    add_textbox(
        slide,
        Inches(5.2),
        Inches(7.12),
        Inches(3),
        Inches(0.28),
        "baltan.xyz",
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        Inches(11.2),
        Inches(7.12),
        Inches(1.5),
        Inches(0.28),
        f"{n} / {TOTAL}",
        size=11,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def section_label(slide, text):
    add_textbox(
        slide,
        Inches(0.7),
        Inches(0.32),
        Inches(8),
        Inches(0.28),
        text.upper(),
        size=12,
        bold=True,
        color=ACCENT,
    )


def title(slide, text, size=30):
    add_textbox(slide, Inches(0.7), Inches(0.55), Inches(12), Inches(0.65), text, size=size, bold=True)


def subtitle(slide, text):
    add_textbox(slide, Inches(0.7), Inches(1.2), Inches(12), Inches(0.4), text, size=15, color=MUTED)


def bullet_block(slide, left, top, width, height, items, size=15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        r = p.add_run()
        set_run(r, f"●  {item}", size=size, color=TEXT)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_logo(slide, left, top, height=Inches(0.5)):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), left, top, height=height)


def add_icon(slide, left, top, size=Inches(0.55)):
    if ICON.exists():
        slide.shapes.add_picture(str(ICON), left, top, width=size, height=size)


def card(slide, left, top, width, height, headline, body, accent_top=False):
    rect(slide, left, top, width, height, fill=ELEV, line=LINE)
    if accent_top:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(4))
        fill_shape(bar, ACCENT)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.2), width - Inches(0.4), Inches(0.4), headline, size=15, bold=True, color=ACCENT)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.65), width - Inches(0.4), height - Inches(0.85), body, size=13, color=TEXT)


# ─── Slides ───────────────────────────────────────────────────────────────


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    # left accent strip
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.14), SLIDE_H)
    fill_shape(bar, ACCENT)

    add_logo(s, Inches(0.75), Inches(1.35), Inches(0.9))
    chip(s, Inches(0.75), Inches(2.5), Inches(3.4), Inches(0.38), "PRARAMBH 2026  ·  Closed pilot")

    add_textbox(
        s,
        Inches(0.75),
        Inches(3.15),
        Inches(11.5),
        Inches(1.2),
        "A health check for the database\nthat keeps your apps fast.",
        size=32,
        bold=True,
    )
    add_textbox(
        s,
        Inches(0.75),
        Inches(4.5),
        Inches(11),
        Inches(0.55),
        "Explainable Redis / Valkey health — without opening it to the internet.",
        size=16,
        color=MUTED,
    )
    add_textbox(
        s,
        Inches(0.75),
        Inches(5.5),
        Inches(9),
        Inches(0.9),
        "Delhi  ·  tandonkartikey11@gmail.com  ·  +91 8077511218\nbaltan.xyz",
        size=14,
        color=MUTED,
    )
    notes(
        s,
        "Baltan is a simple idea: a health check for the fast database behind modern apps — "
        "without opening that database to the internet. Pitching at PRARAMBH 2026.",
    )


def slide_team(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 2)
    section_label(s, "Team")
    title(s, "One founder. Full ownership.")
    subtitle(s, "Built the product. Talks to every customer. Owns every decision.")

    rect(s, Inches(0.7), Inches(1.9), Inches(12), Inches(4.5), fill=ELEV, line=LINE)
    accent_bar(s, Inches(0.7), Inches(1.9), Inches(0.1), Inches(4.5))
    add_icon(s, Inches(1.1), Inches(2.25), Inches(0.85))
    add_textbox(s, Inches(2.2), Inches(2.25), Inches(9), Inches(0.45), "Kartikey Tandon", size=28, bold=True)
    add_textbox(
        s,
        Inches(2.2),
        Inches(2.8),
        Inches(9),
        Inches(0.35),
        "Founder & CEO  ·  Solo co-founder  ·  Equity 100%",
        size=15,
        color=ACCENT,
    )

    cols = [
        ("What I do", "Build the product, run demos, talk to startups & DevOps teams, and steer go-to-market."),
        ("Why me", "Shipped Baltan end-to-end — monitor, dashboard, alerts — and already running real pilots."),
        ("Proof", "100+ demos booked. Pilot interest from CARPL AI, BigOHealth, and DevOps practitioners."),
    ]
    x = 1.1
    for h, b in cols:
        add_textbox(s, Inches(x), Inches(3.55), Inches(3.5), Inches(0.35), h, size=13, bold=True, color=MUTED)
        add_textbox(s, Inches(x), Inches(4.0), Inches(3.5), Inches(1.6), b, size=14)
        x += 3.8

    add_textbox(
        s,
        Inches(1.1),
        Inches(5.75),
        Inches(11),
        Inches(0.35),
        "After incubation support: hire sales + infrastructure help — keep product ownership with founder.",
        size=13,
        color=MUTED,
    )
    notes(s, "Solo founder, 100% equity. Built everything, 100+ demos. Will hire after grant support.")


def slide_problem(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 3)
    section_label(s, "Problem")
    title(s, "Apps feel slow. Teams don’t know why.")
    subtitle(
        s,
        "Most modern apps use a super-fast memory database (Redis / Valkey). When it gets sick, the whole product feels broken.",
    )

    pains = [
        ("Memory fills up", "Old data never expires. The database grows until the app crashes or slows down."),
        ("Hidden bottlenecks", "One oversized piece of data can freeze everyone else — like one truck blocking a highway."),
        ("Alerts without answers", "Tools say “database is slow” but not what to fix today."),
        ("Security says no", "Many tools need the database opened to the internet. Security teams refuse."),
    ]
    x = 0.7
    for h, b in pains:
        rect(s, Inches(x), Inches(1.95), Inches(2.95), Inches(2.55), fill=ELEV, line=LINE)
        top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.95), Inches(2.95), Pt(5))
        fill_shape(top, ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(2.25), Inches(2.6), Inches(0.55), h, size=15, bold=True, color=ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(2.9), Inches(2.6), Inches(1.4), b, size=13)
        x += 3.1

    rect(s, Inches(0.7), Inches(4.75), Inches(12), Inches(1.85), fill=ELEV2, line=LINE)
    add_textbox(s, Inches(0.95), Inches(4.95), Inches(11.5), Inches(0.35), "Today’s workarounds", size=13, bold=True, color=MUTED)
    add_textbox(
        s,
        Inches(0.95),
        Inches(5.4),
        Inches(11.5),
        Inches(0.9),
        "Browse data manually  ·  Watch generic graphs  ·  Ask a senior engineer who “just knows”\n"
        "Baltan is different: plain-English findings + a clear health score — and the database stays private.",
        size=15,
    )
    notes(
        s,
        "Explain Redis simply as the fast memory database. Pain: crashes, mystery slowdowns, no clear fix, security blocks. "
        "30 seconds.",
    )


def slide_product(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 4)
    section_label(s, "Product")
    title(s, "Baltan = doctor for your fast database")
    subtitle(s, "We host a simple dashboard. You place a small safe monitor next to your database. Nothing private leaves.")

    steps = [
        ("1", "Database stays private", "Never opened to Baltan or the public internet."),
        ("2", "Safe check-ups", "Reads health signals only — never your customer data or passwords."),
        ("3", "Clear answers", "A score out of 100 + “what’s wrong” and “what to do next.”"),
    ]
    x = 0.7
    for num, h, b in steps:
        rect(s, Inches(x), Inches(1.9), Inches(3.85), Inches(2.15), fill=ELEV, line=LINE)
        circ = oval(s, Inches(x + 0.25), Inches(2.15), Inches(0.45), Inches(0.45), ACCENT_DIM)
        # number over circle area
        add_textbox(s, Inches(x + 0.25), Inches(2.2), Inches(0.45), Inches(0.4), num, size=16, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(x + 0.85), Inches(2.2), Inches(2.7), Inches(0.4), h, size=16, bold=True)
        add_textbox(s, Inches(x + 0.25), Inches(2.85), Inches(3.4), Inches(0.95), b, size=13, color=MUTED)
        x += 4.05

    feats = [
        ("Health score", "See at a glance if things are healthy, okay, or urgent."),
        ("Plain-English findings", "Missing expiry, oversized data, risky commands — explained simply."),
        ("Alerts on Slack", "The team hears about problems early — not after customers complain."),
        ("One-click install", "A single small monitor beside the database. Minutes, not weeks."),
    ]
    x = 0.7
    for h, b in feats:
        rect(s, Inches(x), Inches(4.3), Inches(2.95), Inches(2.2), fill=ELEV, line=LINE)
        add_textbox(s, Inches(x + 0.18), Inches(4.5), Inches(2.6), Inches(0.45), h, size=14, bold=True, color=ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(5.05), Inches(2.6), Inches(1.2), b, size=13)
        x += 3.1

    notes(
        s,
        "Doctor metaphor. Three steps: private, safe check-ups, clear answers. Features without jargon. 30 seconds.",
    )


def slide_compete(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 5)
    section_label(s, "Why Baltan wins")
    title(s, "Others show graphs. We explain the problem.")
    subtitle(s, "Built for teams who care about privacy and need clear next steps.")

    headers = ["", "Do it yourself", "Browse tools", "Big monitoring tools", "Baltan"]
    rows = [
        ["Explains what’s wrong", "No", "Rarely", "Sometimes", "Yes"],
        ["Says what to fix next", "Depends on expert", "Explore yourself", "Generic alerts", "Yes"],
        ["Keeps database private", "Yes", "Often local only", "Varies", "Always"],
        ["Easy for a small team", "Hard", "Medium", "Heavy & costly", "Yes"],
        ["Made for Redis / Valkey", "No", "Browse-focused", "One of many", "Yes"],
    ]
    col_w = [3.0, 2.1, 2.2, 2.5, 2.2]
    y = 1.85
    rect(s, Inches(0.7), Inches(y), Inches(12), Inches(0.48), fill=ACCENT_DIM, line=LINE)
    x = 0.7
    for i, h in enumerate(headers):
        add_textbox(
            s,
            Inches(x + 0.08),
            Inches(y + 0.1),
            Inches(col_w[i] - 0.1),
            Inches(0.35),
            h,
            size=12,
            bold=True,
            color=ACCENT if i == 4 else MUTED,
        )
        x += col_w[i]

    y = 2.4
    for ri, row in enumerate(rows):
        bg = ELEV if ri % 2 == 0 else ELEV2
        rect(s, Inches(0.7), Inches(y), Inches(12), Inches(0.55), fill=bg, line=LINE)
        x = 0.7
        for i, cell in enumerate(row):
            add_textbox(
                s,
                Inches(x + 0.08),
                Inches(y + 0.12),
                Inches(col_w[i] - 0.1),
                Inches(0.35),
                cell,
                size=12,
                bold=(i == 0 or i == 4),
                color=ACCENT if i == 4 else TEXT,
            )
            x += col_w[i]
        y += 0.55

    add_textbox(
        s,
        Inches(0.7),
        Inches(5.4),
        Inches(12),
        Inches(1.1),
        "Strategy: win teams that refuse to expose their database — then become the trusted Redis health layer inside the company.",
        size=15,
        color=MUTED,
    )
    notes(s, "Simple comparison. We explain + protect privacy + stay easy. 25 seconds.")


def slide_business(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 6)
    section_label(s, "Business model")
    title(s, "Start free. Grow into a paid cloud plan.")
    subtitle(s, "Who pays: engineering / DevOps teams. Early pricing kept simple for Indian startups & global pilots.")

    items = [
        ("Who pays?", "Companies that run apps on Redis / Valkey — startups to mid-size tech teams."),
        ("What they buy", "A hosted health dashboard + alerts. They keep the small monitor on their side."),
        ("How we sell", "Demo → free pilot → paid plan when they see value."),
        ("Revenue model", "Monthly subscription (SaaS). Charge per database / team plan."),
    ]
    x = 0.7
    for h, b in items:
        rect(s, Inches(x), Inches(1.85), Inches(2.95), Inches(1.85), fill=ELEV, line=LINE)
        add_textbox(s, Inches(x + 0.18), Inches(2.05), Inches(2.6), Inches(0.35), h, size=13, bold=True, color=ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(2.5), Inches(2.6), Inches(1.0), b, size=13)
        x += 3.1

    add_textbox(s, Inches(0.7), Inches(3.9), Inches(12), Inches(0.35), "Proposed plans (to be finalized after pilots)", size=13, bold=True, color=MUTED)

    tiers = [
        ("Free Pilot", "₹0", "Basic health + findings\nTime-boxed · limited databases"),
        ("Starter", "₹2,999/mo\n(~$35)", "Up to 3 databases\nSlack alerts · core reports"),
        ("Growth", "₹7,999/mo\n(~$95)", "Up to 10 databases\nLonger history · digests"),
        ("Scale", "Custom", "More databases\nOnboarding · priority support"),
    ]
    x = 0.7
    for name, price, desc in tiers:
        highlight = name == "Growth"
        rect(s, Inches(x), Inches(4.35), Inches(2.95), Inches(2.2), fill=ELEV, line=ACCENT if highlight else LINE)
        add_textbox(s, Inches(x + 0.18), Inches(4.5), Inches(2.6), Inches(0.35), name, size=15, bold=True)
        add_textbox(s, Inches(x + 0.18), Inches(4.95), Inches(2.6), Inches(0.65), price, size=14, bold=True, color=ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(5.65), Inches(2.6), Inches(0.75), desc, size=12, color=MUTED)
        x += 3.1

    notes(
        s,
        "Free pilot then paid SaaS. Modest India-first pricing. No revenue yet — honest. 30 seconds.",
    )


def slide_market(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 7)
    section_label(s, "Market opportunity")
    title(s, "Every modern app needs a healthy fast database")
    subtitle(s, "Global need. India beachhead. Sell where trust and clarity matter.")

    cards = [
        ("Who we sell to", "DevOps, platform, and product engineering teams — plus freelancers who manage many apps."),
        ("Where", "Global product. Start with India (Delhi / NCR) for hands-on pilots, then remote US & Europe."),
        ("Willingness to pay", "Strong interest already: CARPL AI, BigOHealth, DevOps individuals. 100+ demos. Not paid yet."),
        ("Near-term opportunity", "Illustrative: hundreds of fit teams × a few databases × affordable monthly plans. Focus on Redis health first."),
    ]
    positions = [(0.7, 1.85), (6.85, 1.85), (0.7, 4.15), (6.85, 4.15)]
    for (h, b), (lx, ty) in zip(cards, positions):
        card(s, Inches(lx), Inches(ty), Inches(5.8), Inches(2.05), h, b, accent_top=True)

    notes(
        s,
        "ICP plain language. India beachhead. Named interest + 100+ demos. Keep market estimate humble. 25 seconds.",
    )


def slide_status(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 8)
    section_label(s, "Present status")
    title(s, "Working product. Early customers. Zero revenue.")
    subtitle(s, "Honest stage: closed pilot with real demos and design partners — monetization next.")

    chip(s, Inches(0.7), Inches(1.75), Inches(4.2), Inches(0.4), "STAGE: Early customers / Pilot")

    rect(s, Inches(0.7), Inches(2.4), Inches(5.9), Inches(4.1), fill=ELEV, line=LINE)
    add_textbox(s, Inches(0.95), Inches(2.6), Inches(5.4), Inches(0.4), "What exists today", size=15, bold=True, color=ACCENT)
    bullet_block(
        s,
        Inches(0.95),
        Inches(3.15),
        Inches(5.4),
        Inches(3.1),
        [
            "Live product website: baltan.xyz",
            "Safe monitor + hosted health dashboard",
            "Plain-English findings & health score",
            "Slack alerts · multi-database support",
            "Install in minutes for pilot teams",
        ],
        size=14,
    )

    rect(s, Inches(6.85), Inches(2.4), Inches(5.9), Inches(4.1), fill=ELEV, line=LINE)
    add_textbox(s, Inches(7.1), Inches(2.6), Inches(5.4), Inches(0.4), "Validation so far", size=15, bold=True, color=ACCENT)
    bullet_block(
        s,
        Inches(7.1),
        Inches(3.15),
        Inches(5.4),
        Inches(3.1),
        [
            "100+ demos booked",
            "Pilot interest: CARPL AI, BigOHealth",
            "DevOps individuals testing the idea",
            "Free basic pilot open",
            "No paying customers yet (by design)",
        ],
        size=14,
    )
    notes(s, "Stage chip. Built vs validation. Honest about $0. 30 seconds.")


def slide_funding(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    footer(s, 9)
    section_label(s, "Funding requirement")
    title(s, "Why we need support — and where it goes")
    subtitle(
        s,
        "Incubation support helps turn a working pilot into a reliable, paid product for more teams.",
    )

    # Why needed
    whys = [
        ("Product maturity", "Harden the monitor and dashboard so pilots run smoothly for months, not days."),
        ("Trust & reliability", "Stronger hosting, backups, and security basics so companies feel safe to depend on Baltan."),
        ("Reach customers", "More demos, onboarding help, and early sales support to convert interest into paying teams."),
        ("Founder focus", "Cover core operating needs so time stays on building and talking to customers — not side work."),
    ]
    x = 0.7
    for h, b in whys:
        rect(s, Inches(x), Inches(1.85), Inches(2.95), Inches(2.2), fill=ELEV, line=LINE)
        top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.85), Inches(2.95), Pt(5))
        fill_shape(top, ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(2.15), Inches(2.6), Inches(0.45), h, size=14, bold=True, color=ACCENT)
        add_textbox(s, Inches(x + 0.18), Inches(2.7), Inches(2.6), Inches(1.15), b, size=13)
        x += 3.1

    # Where used + goals
    rect(s, Inches(0.7), Inches(4.3), Inches(6.0), Inches(2.3), fill=ELEV, line=LINE)
    add_textbox(s, Inches(0.95), Inches(4.45), Inches(5.5), Inches(0.35), "Where funds will be used", size=14, bold=True, color=ACCENT)
    bullet_block(
        s,
        Inches(0.95),
        Inches(4.9),
        Inches(5.5),
        Inches(1.55),
        [
            "Product reliability & usability polish",
            "Cloud hosting & security basics",
            "Founder-led pilots, demos & onboarding",
            "Legal, compliance & day-to-day ops",
        ],
        size=13,
    )

    rect(s, Inches(6.95), Inches(4.3), Inches(5.8), Inches(2.3), fill=ELEV, line=LINE)
    add_textbox(s, Inches(7.2), Inches(4.45), Inches(5.3), Inches(0.35), "What success looks like", size=14, bold=True, color=ACCENT)
    bullet_block(
        s,
        Inches(7.2),
        Inches(4.9),
        Inches(5.3),
        Inches(1.55),
        [
            "More active pilot teams running smoothly",
            "First paid conversions",
            "Pricing finalized from real usage",
            "Ready for the next growth stage",
        ],
        size=13,
    )

    notes(
        s,
        "Do not quote rupee amounts. Explain why support is needed and where money goes: product, trust, customers, ops. "
        "Close with success outcomes and invite Q&A.",
    )


def write_companion_docs():
    (OUT_DIR / "SPEAKER_SCRIPT_4MIN.md").write_text(
        """# Baltan — 4-minute script (PRARAMBH / non-technical friendly)

### Slide 1 — Cover (~20s)
Baltan is a health check for the fast database that keeps apps running — without opening that database to the internet. Pitching at PRARAMBH 2026.

### Slide 2 — Team (~25s)
I’m Kartikey, solo founder, 100% equity. I built the product and run every demo — 100+ already. After incubation support I’ll add sales help.

### Slide 3 — Problem (~30s)
When the fast database gets sick, apps feel slow or crash. Teams get alerts without answers. Security often blocks tools that need an open door to the internet.

### Slide 4 — Product (~30s)
Think of Baltan as a doctor: the database stays private, we run safe check-ups, and we return a score plus plain-English “what to do next.” Slack alerts. Installs in minutes.

### Slide 5 — Why we win (~25s)
Browse tools and big monitors don’t explain Redis problems clearly. We do — and we never need the database opened to us.

### Slide 6 — Business (~30s)
Free pilot, then a simple monthly plan. Proposed from about ₹2,999/month. No revenue yet — we’re converting interest into paid.

### Slide 7 — Market (~25s)
Global need; start in India. CARPL AI, BigOHealth, DevOps individuals already interested. 100+ demos.

### Slide 8 — Status (~30s)
Working product at baltan.xyz. Early customers / closed pilot. Free pilot live. Zero paying customers so far — by design while we learn.

### Slide 9 — Funding (~30s)
We need incubation support to mature the product, earn trust on security and reliability, and convert pilots into paying teams. Funds go to product polish, hosting & security, demos/onboarding, and basic ops — so the founder stays focused on customers. Happy to take questions.

---

# Q&A (short)

1. **What is Redis?** The fast memory database many apps use so pages and APIs feel instant.
2. **Is data safe?** Yes — we never read customer contents or passwords; database stays closed to us.
3. **Why no revenue?** Free pilots first; pricing after proof. Interest is already strong.
4. **Solo founder?** Yes today; support enables first hire for sales/support when ready.
5. **How much funding?** We’re focused on why and use of funds; happy to discuss fit with incubator programs in Q&A.
6. **vs Datadog / big tools?** They watch many things; we explain this one critical database clearly and affordably.
7. **India vs global?** Build trust in India first; product works globally.
8. **Next stage after this?** Paid conversions and a clear path to the next growth stage once pilots prove value.
""",
        encoding="utf-8",
    )


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_cover(prs)
    slide_team(prs)
    slide_problem(prs)
    slide_product(prs)
    slide_compete(prs)
    slide_business(prs)
    slide_market(prs)
    slide_status(prs)
    slide_funding(prs)

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    write_companion_docs()
    print(f"Wrote {OUT_PPTX}")
    print(f"Wrote {OUT_DIR / 'SPEAKER_SCRIPT_4MIN.md'}")


if __name__ == "__main__":
    main()
